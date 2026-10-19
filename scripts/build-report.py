"""Build the academic report from the retained Word template and real evidence.
Run using the Codex bundled Python (python-docx 1.2.0 and Pillow 12.3.0).
Rendering/visual QA is a separate mandatory step after this builder.
"""
from pathlib import Path
from copy import deepcopy
import hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent.parent
TEMPLATE=ROOT/'docs/report-template.docx'
assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()=='02c9189f6e9a969c92e76b214a255d133b5261f8e1f9b9198b68a3f6b6149a13'
doc=Document(TEMPLATE)
doc.core_properties.author='Nguyễn Xuân Hải'
doc.core_properties.last_modified_by='Nguyễn Xuân Hải'
doc.core_properties.title='Triển khai hệ thống quản lý tài sản và thiết bị Hải Asset'
doc.core_properties.subject='Đề 16 Triển khai và Quản trị Hệ thống Phần mềm'
doc.core_properties.keywords='dtc245200551, CNTT K23D, WSL2, Docker Engine, PostgreSQL'

# Preserve the source section geometry and the first eight cover paragraphs.
cover=doc.paragraphs[:8]
for element in list(doc.element.body):
    if element.tag==qn('w:sectPr') or element in [p._p for p in cover]: continue
    doc.element.body.remove(element)
for style in doc.styles:
    if style.type==1:
        style.font.name='Times New Roman'
        style.font.color.rgb=RGBColor(0,0,0)
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for fonts in style.element.iter(qn('w:rFonts')):
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            fonts.attrib.pop(qn('w:'+key),None)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(13)
normal.paragraph_format.space_before=Pt(0);normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.15
for name,size in [('Title',16),('Heading 1',14),('Heading 2',13)]:
    s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.font.bold=True;s.paragraph_format.space_before=Pt(0 if name!='Heading 2' else 10)
    s.paragraph_format.space_after=Pt(9);s.paragraph_format.keep_with_next=True
    s.paragraph_format.line_spacing=1.1
doc.styles['Heading 2'].font.italic=True
for name in ['Caption','Footer']:
    doc.styles[name].font.name='Times New Roman';doc.styles[name].font.size=Pt(10)
    doc.styles[name].font.color.rgb=RGBColor(0,0,0)

def cover_text(p,text,size=13,bold=False,italic=False,align=WD_ALIGN_PARAGRAPH.CENTER,before=0,after=0):
    p.clear();p.alignment=align
    for border in list(p._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
    p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.line_spacing=1.15
    r=p.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(size)
    r.bold=bold;r.italic=italic;r.font.color.rgb=RGBColor(0,0,0)

cover_text(cover[0],'TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG\nKHOA CÔNG NGHỆ THÔNG TIN',12,True,after=24)
cover_text(cover[1],'',after=10)
cover[2].style='Title'
cover_text(cover[2],'BÁO CÁO BÀI TẬP LỚN\nTRIỂN KHAI VÀ QUẢN TRỊ\nHỆ THỐNG PHẦN MỀM',16,True,after=28)
cover_text(cover[3],'ĐỀ 16',14,True,after=10)
cover_text(cover[4],'HỆ THỐNG QUẢN LÝ TÀI SẢN VÀ THIẾT BỊ\nHẢI ASSET',16,True,after=48)
cover_text(cover[5],'Sinh viên thực hiện: Nguyễn Xuân Hải\nMã số sinh viên: dtc245200551\nLớp: CNTT K23D\nGiảng viên hướng dẫn: Vũ Việt Dũng',13,italic=True,align=WD_ALIGN_PARAGRAPH.RIGHT,after=32)
cover_text(cover[6],'',after=14)
cover_text(cover[7],'Thái Nguyên, tháng 10 năm 2026',12,True)

section=doc.sections[0]
section.different_first_page_header_footer=True
fp=section.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=fp.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)

page_titles=[]
def page(title):
    doc.add_page_break();doc.add_paragraph(title,'Heading 1');page_titles.append(title)
def p(text):
    out=doc.add_paragraph(text);out.paragraph_format.widow_control=True;return out
def h(text): return doc.add_paragraph(text,'Heading 2')
def code(text):
    out=doc.add_paragraph();out.paragraph_format.space_after=Pt(8)
    out.paragraph_format.line_spacing=1.0
    r=out.add_run(text);r.font.name='Consolas';r.font.size=Pt(9.5)
    shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'F2F4F7');out._p.get_or_add_pPr().append(shading)
    return out
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,width in enumerate(widths):t.columns[i].width=Inches(width)
    for i,text in enumerate(headers):t.rows[0].cells[i].text=str(text)
    for row in rows:
        cells=t.add_row().cells
        for i,text in enumerate(row):cells[i].text=str(text)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders);margins=OxmlElement('w:tcMar')
            for side,value in [('top',85),('bottom',85),('left',110),('right',110)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'253C5B' if ri==0 else ('F1F5FA' if ri%2==0 else 'FFFFFF'));pr.append(shading)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.05
                for r in para.runs:
                    r.font.name='Times New Roman';r.font.size=Pt(11);r.bold=ri==0
                    r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
        if ri==0:
            e=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(e)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def figure(name,caption,width=6.1,diagram=False):
    source=ROOT/'docs'/('diagrams' if diagram else 'evidence')/name
    assert source.exists(),source
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before=Pt(4);para.paragraph_format.space_after=Pt(4)
    para.paragraph_format.keep_with_next=True
    shape=para.add_run().add_picture(str(source),width=Inches(width))
    shape._inline.docPr.set('descr',caption)
    cap=doc.add_paragraph(caption,'Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after=Pt(9)

# Page 2
page('MỤC LỤC')
contents=[
('Chương 1 Tổng quan kiến trúc và môi trường',3),('1.1 Mục tiêu và phạm vi',3),('1.2 Kiến trúc triển khai WSL2',4),('1.3 Quản lý mã nguồn trên GitHub',5),
('Chương 2 Ứng dụng Database và Reverse Proxy',6),('2.1 Mô hình dữ liệu',6),('2.2 Hồ sơ thiết bị và tra cứu',7),('2.3 Vòng đời bảo trì và transaction',8),('2.4 PostgreSQL và pgAdmin',9),('2.5 Nginx và cách khởi động',10),
('Chương 3 Giám sát và log tập trung',11),('3.1 Prometheus và Grafana',11),('3.2 Luồng thu thập log',12),('3.3 Truy vấn LogQL',13),
('Chương 4 Hardening và đánh giá',14),('4.1 Biện pháp gia cố',14),('4.2 Kết quả kiểm thử',15),('4.3 Sao lưu và phục hồi',16),('Kết luận và kịch bản demo',17),('Tài liệu tham khảo',18)]
for label,n in contents:
    para=doc.add_paragraph();para.paragraph_format.space_after=Pt(6)
    para.paragraph_format.tab_stops.add_tab_stop(Inches(6.0),WD_ALIGN_PARAGRAPH.RIGHT)
    r=para.add_run(label);r.bold=label.startswith(('Chương','Kết luận','Tài liệu'));para.add_run('\t'+str(n))
p('Repository: dtc245200551-jpg/dtc245200551-de16-quan-ly-tai-san. Nhánh bàn giao: codex/de16-quan-ly-tai-san. Các hình là minh chứng của hệ thống chạy bằng Docker Engine trong Ubuntu WSL2.')

# Page 3
page('CHƯƠNG 1 TỔNG QUAN KIẾN TRÚC VÀ MÔI TRƯỜNG')
h('1.1 Mục tiêu và phạm vi')
p('Đề tài triển khai website quản lý tài sản và thiết bị Hải Asset, giúp quản trị viên theo dõi từng thiết bị, vị trí sử dụng, trạng thái và lịch sử bảo trì. Bài thực hành tập trung đồng thời vào nghiệp vụ và vận hành: ứng dụng phải kết nối cơ sở dữ liệu, truy cập qua reverse proxy, có số liệu giám sát, log tập trung và biện pháp gia cố được kiểm chứng.')
p('Mỗi hồ sơ tương ứng một thiết bị riêng lẻ, được phân biệt bằng mã và serial. Người quản trị có thể tìm kiếm, lọc danh sách, lập lịch bảo trì, ghi nhận thời điểm bắt đầu, kết quả và chi phí. Lịch sử được giữ để đối chiếu sau khi hoàn tất. Phạm vi một vai trò quản trị giúp thao tác demo rõ ràng; chưa triển khai cấp phát/thu hồi, nhiều đơn vị độc lập hoặc kế toán khấu hao.')
h('Môi trường thực hiện')
table(['Thành phần','Cấu hình đã sử dụng'],[
('Máy người dùng','Windows, Chrome và WSL2'),('Hệ điều hành chạy dịch vụ','Ubuntu 24.04.5 LTS, filesystem Linux'),('Runtime container','Docker Engine 29.9.1; Docker Compose 5.6.0; cgroup v2'),('Ứng dụng','Node.js 24.22.0; Express 5.1.0; EJS'),('Dữ liệu','PostgreSQL 17.11; pgAdmin 9.15'),('Nguồn chính','/home/hai/projects/de16-quan-ly-tai-san')],[1.7,4.5])
p('Docker Engine được cài từ APT repository chính thức của Docker và được systemd quản lý trong Ubuntu. Các lệnh vận hành chọn socket unix:///var/run/docker.sock. Nguồn chạy nằm trong filesystem Linux để tránh nhầm quyền file và mount khi đặt trên ổ Windows. Trình duyệt Windows truy cập dịch vụ qua localhost được WSL2 chuyển tiếp.')
p('Dữ liệu khởi tạo có 12 thiết bị, ba danh mục, ba vị trí và ba phiếu bảo trì. Các script có thể chạy lại mà không tạo bản ghi trùng hoặc đặt lại mật khẩu quản trị đã có.')

# Page 4
page('1.2 Kiến trúc triển khai WSL2')
p('Hệ thống là một ứng dụng web nguyên khối được chia thành các dịch vụ triển khai container. Nginx nhận HTTP, chuyển yêu cầu vào Express; ứng dụng render EJS và truy vấn PostgreSQL. pgAdmin là công cụ quản lý DB riêng. Các thành phần giám sát/log được trình bày tại Chương 3.')
figure('architecture.png','Hình 1.1 Kiến trúc triển khai ứng dụng trong Ubuntu WSL2',4.6,True)
p('Nginx và app giao tiếp qua mạng frontend; app và PostgreSQL dùng mạng database nội bộ. Database không publish cổng 5432 ra Windows. Dữ liệu PostgreSQL được lưu trong volume postgres_data và tồn tại sau khi dừng container.')

# Page 5
page('1.3 Quản lý mã nguồn trên GitHub')
p('Repository private thuộc tài khoản dtc245200551-jpg, dùng nhánh codex/de16-quan-ly-tai-san làm nhánh mặc định. Mã nguồn, schema, Compose, Nginx, Prometheus, Grafana, Loki và Promtail đều nằm trong cùng repository để có thể dựng lại hệ thống. Danh tính commit là Nguyễn Xuân Hải, email trường gắn với tài khoản GitHub.')
table(['Mốc','Nội dung có thể kiểm chứng'],[
('1','Ứng dụng, PostgreSQL, pgAdmin và Nginx'),('2','Prometheus, Grafana, exporter và dashboard'),('3','Loki, Promtail, logrotate và LogQL'),('4','Hardening, kiểm thử, backup/restore'),('5','Báo cáo, sơ đồ, ảnh và kịch bản demo')],[0.6,5.6])
figure('github-history.png','Hình 1.2 Bốn mốc kỹ thuật trên GitHub trước khi đóng gói báo cáo',6.1)
p('Các mốc được commit theo thời gian thực và push sau từng lần, cách nhau ít nhất mười phút. Commit bàn giao tài liệu bổ sung chính báo cáo này. Lịch sử trực tiếp trên GitHub là nơi tra cứu đầy đủ mã commit và thời điểm cuối cùng.')
p('Thư mục .local chứa mật khẩu, kết quả chạy và backup nên bị Git bỏ qua. Các file cấu hình công khai chỉ tham chiếu đường dẫn secret. Repository private cần được chủ tài khoản cấp quyền đọc cho giảng viên khi nộp bài.')

# Page 6
page('CHƯƠNG 2 ỨNG DỤNG DATABASE VÀ REVERSE PROXY')
h('2.1 Mô hình dữ liệu')
p('Bốn bảng nghiệp vụ dùng khóa ngoại để liên kết thiết bị với danh mục, vị trí và các lần bảo trì. Một danh mục/vị trí có thể chưa có hoặc có nhiều thiết bị; một thiết bị có thể có nhiều phiếu bảo trì theo thời gian.')
figure('erd.png','Hình 2.1 ERD logic của bốn bảng nghiệp vụ theo schema SQL',4.7,True)
p('Hai bảng kỹ thuật admins và sessions lưu tài khoản và phiên đăng nhập. Sessions lưu JSON, không có khóa ngoại trực tiếp tới admins. Unique index riêng giới hạn một phiếu in_progress trên mỗi tài sản; tổng số phiếu lịch sử không bị giới hạn bởi ràng buộc này.')

# Page 7
page('2.2 Hồ sơ thiết bị và tra cứu')
p('Trang Thiết bị tập hợp mã, tên, serial, loại, vị trí, giá trị và trạng thái. Danh sách sắp theo mã và chia mười hồ sơ mỗi trang. Người quản trị tìm theo mã/tên/serial hoặc kết hợp bộ lọc trạng thái, danh mục và vị trí. Khi chuyển trang, các điều kiện lọc được giữ lại.')
figure('assets.png','Hình 2.2 Danh sách thiết bị thực tế và các bộ lọc',6.1)
table(['Dữ liệu','Quy tắc lưu trữ'],[
('Mã thiết bị','Bắt buộc, duy nhất; chuẩn hóa chữ hoa'),('Serial','Không bắt buộc; nếu có phải duy nhất'),('Danh mục và vị trí','Khóa ngoại bắt buộc, tham chiếu bản ghi tồn tại'),('Ngày mua và giá trị','Ngày hợp lệ; VND nguyên không âm'),('Trạng thái','available, in_use, maintenance hoặc retired')],[1.7,4.5])
p('Cập nhật hồ sơ dùng trường version. Nếu bản ghi đã thay đổi ở phiên khác, server trả 409 và yêu cầu tải lại để tránh ghi đè dữ liệu mới. Dữ liệu được kiểm tra tại server; các truy vấn tìm kiếm dùng tham số thay vì ghép trực tiếp từ khóa vào SQL.')
p('Giao diện dùng EJS để escape nội dung hiển thị. Các form và bảng có quy tắc CSS cho màn hình hẹp; bảng dài cuộn trong vùng danh sách, còn form chuyển về một cột. Thông báo tiếng Việt nêu nguyên nhân dữ liệu chưa hợp lệ để người dùng kiểm tra lại.')

# Page 8
page('2.3 Vòng đời bảo trì và transaction')
p('Phiếu bảo trì đi từ đã lên lịch tới đang thực hiện, sau đó hoàn tất hoặc hủy. Khi bắt đầu, hệ thống ghi trạng thái trước đó của thiết bị và chuyển thiết bị sang maintenance. Hoàn tất lưu kết quả, chi phí, thời điểm kết thúc rồi khôi phục trạng thái trước. Hủy phiếu đang thực hiện cũng khôi phục trạng thái và giữ lý do hủy.')
figure('asset-history.png','Hình 2.3 TB-002 có lịch sử hoàn tất với chi phí 150.000 VND',6.1)
table(['Tình huống','Xử lý'],[
('Hai yêu cầu bắt đầu cùng lúc','Khóa dòng thiết bị; chỉ một yêu cầu thành công'),('Đổi tay trạng thái khi bảo trì','Từ chối; cần kết thúc hoặc hủy phiếu'),('Ngừng sử dụng khi còn lịch mở','Từ chối để xử lý lịch còn tồn tại'),('Ghi lại phiếu đã hoàn tất','Từ chối; giữ lịch sử đã kết thúc')],[2.4,3.8])
p('Mỗi thao tác khóa thiết bị trước rồi khóa phiếu bảo trì theo cùng một thứ tự, thực hiện thay đổi trong transaction. Nếu có lỗi, transaction rollback để tránh thiết bị và phiếu lệch trạng thái. Partial unique index ở PostgreSQL bổ sung rào chắn cho trường hợp nhiều phiếu đồng thời trên một thiết bị.')
p('Dữ liệu minh chứng có TB-002 hoàn tất bảo trì và TB-006 đang bảo trì. Các bước demo mới dùng mã thiết bị riêng để có thể lặp lại mà không sửa lịch sử đã ghi nhận.')

# Page 9
page('2.4 PostgreSQL và pgAdmin')
p('PostgreSQL 17 lưu dữ liệu trong volume riêng của project de16hai. Healthcheck dùng kết nối TCP bằng role ứng dụng và câu SELECT 1; chỉ sau khi database sẵn sàng, migration mới chạy. Migration có advisory lock, tạo schema và cấp quyền, sau đó ứng dụng được khởi động.')
figure('pgadmin.png','Hình 2.4 pgAdmin kết nối server Hải Asset PostgreSQL',6.1)
table(['Role','Trách nhiệm và quyền'],[
('asset_owner','Khởi tạo cluster/role và thực hiện backup/restore'),('asset_migrate','Sở hữu bảng, migration và quản trị dữ liệu demo'),('asset_app','Đọc/ghi nghiệp vụ, quản lý phiên; không xóa hồ sơ hoặc tạo bảng'),('asset_exporter','pg_monitor phục vụ PostgreSQL exporter')],[1.5,4.7])
p('pgAdmin chạy ở localhost:8087 và đã khai báo server postgres:5432, database assets. Tài khoản giao diện là admin@hai.edu.vn. Mật khẩu giao diện và mật khẩu role kết nối là hai secret riêng; báo cáo và ảnh minh chứng chỉ ghi đường dẫn, không ghi giá trị mật khẩu.')
p('Bảng admins lưu mật khẩu băm bằng scrypt cùng auth_version. Sessions lưu phiên trong PostgreSQL để restart app không làm mất toàn bộ phiên. Khi đổi mật khẩu, auth_version tăng; các phiên có phiên bản cũ không còn truy cập được. Seed chỉ thêm bản ghi còn thiếu, giữ nguyên mật khẩu và trạng thái nghiệp vụ đã có.')

# Page 10
page('2.5 Nginx và cách khởi động')
p('Website được truy cập qua Nginx tại localhost:8086. Nginx chuyển tiếp vào app:3000, đặt X-Forwarded-For từ địa chỉ kết nối thực và sử dụng DNS nội bộ Docker để tiếp tục định tuyến khi app được tạo lại. Endpoint stub_status chỉ phục vụ exporter trong mạng giám sát; /metrics và /readyz bị chặn ở cổng website.')
code("upstream hai_app {\n  zone hai_app 64k;\n  server app:3000 resolve;\n}\nlocation / { proxy_pass http://hai_app; }\nlocation = /metrics { return 404; }")
table(['Header','Mục đích'],[
('Content-Security-Policy','Giới hạn nguồn tài nguyên, form và việc nhúng trang'),('X-Content-Type-Options','nosniff để tránh đoán lại kiểu nội dung'),('X-Frame-Options','DENY để chặn nhúng iframe'),('Referrer-Policy','same-origin để hạn chế gửi referrer'),('Permissions-Policy','Tắt quyền camera, microphone, geolocation không cần thiết')],[2.1,4.1])
h('Trình tự chạy trong Ubuntu')
code('cd ~/projects/de16-quan-ly-tai-san\nsh scripts/start.sh\nsh scripts/test.sh\nsh scripts/stop.sh')
p('start.sh chuẩn bị secrets, kiểm tra daemon, validate Compose, build image và chờ dịch vụ sẵn sàng. Nginx được kiểm tra cú pháp rồi reload cấu hình. common.sh chọn các file Compose theo thành phần đang có và luôn dùng socket Engine của Ubuntu. stop.sh dừng project, giữ volume và secret.')
p('Cấu hình này là demo HTTP trên localhost với security headers, đúng lựa chọn được đề cho phép. Khi đưa ra Internet cần triển khai TLS và bật cookie Secure. PowerShell chỉ gọi script qua WSL2; wrapper có tiến trình giữ Ubuntu hoạt động và không shutdown các distro khác.')

# Page 11
page('CHƯƠNG 3 GIÁM SÁT VÀ LOG TẬP TRUNG')
h('3.1 Prometheus và Grafana')
p('Prometheus scrape mỗi 15 giây và có năm jobs: prometheus, hai-app, nginx, postgres và cadvisor. Nginx exporter đọc stub_status; PostgreSQL exporter dùng pg_monitor; cAdvisor lấy số liệu cgroup của container. Trên daemon Ubuntu, kiểm tra xác nhận nhãn project de16hai và service app/nginx/postgres có cả RAM và CPU rate.')
figure('grafana-metrics.png','Hình 3.1 Dashboard có số liệu CPU và RAM thực theo container',6.1)
table(['Nhóm số liệu','Thông tin được theo dõi'],[
('Container','CPU phần trăm một lõi, RAM working set, lưu lượng mạng'),('HTTP','Số request, lỗi 4xx/5xx, độ trễ p95'),('Nginx và PostgreSQL','Trạng thái exporter, kết nối, giao dịch commit/rollback'),('Nghiệp vụ','Số thiết bị theo trạng thái và thao tác bảo trì')],[1.7,4.5])
p('Grafana tự nạp datasource và dashboard từ repository: 15 panel metrics cùng ba panel log. Các biểu đồ rate cần tối thiểu hai lần scrape. Counter HTTP/bảo trì bắt đầu lại khi app restart; gauge tài sản đọc database nên vẫn thể hiện số liệu bền vững. Không dùng số counter làm sổ lịch sử nghiệp vụ.')
p('Prometheus có quy tắc cảnh báo target mất kết nối và exporter không truy cập được PostgreSQL. Bài thực hành xem trạng thái tại /alerts; chưa cấu hình gửi thông báo ra email hoặc Alertmanager.')

# Page 12
page('3.2 Luồng thu thập log')
p('Ứng dụng ghi JSON với thời điểm, mức độ, sự kiện và request_id. Nginx ghi access log JSON vào file thường trên volume; PostgreSQL dùng logging_collector với jsonlog. Promtail mount ba volume ở chế độ chỉ đọc, lưu positions trong volume riêng rồi gửi dữ liệu tới Loki. Grafana truy vấn Loki qua mạng nội bộ.')
figure('observability.png','Hình 3.2 Luồng metrics và log của riêng project de16hai',6.1,True)
table(['Nguồn','Đặc điểm'],[
('App','Sự kiện bảo trì, HTTP và lỗi đã được lọc thông tin nhạy cảm'),('Nginx','Method, URI không có query, status, duration và request_id'),('PostgreSQL','JSON gồm mức độ và thông điệp; không log statement thông thường')],[1.4,4.8])
p('request_id được giữ trong nội dung JSON để truy vết, không trở thành label có số lượng giá trị tăng không giới hạn. Promtail không đọc log toàn bộ máy và không mount Docker socket. App/Nginx dùng timestamp trong JSON; timestamp gốc của PostgreSQL nằm trong bản ghi, còn thời điểm index là lúc Promtail đọc.')
p('Logrotate kiểm tra mỗi phút, xoay ở ngưỡng 5 MB và giữ ba bản nén; file PostgreSQL theo ngày được dọn sau bảy ngày. Loki giữ dữ liệu 168 giờ qua compactor. Đã thử xoay log thật và kiểm tra các file gzip của app, Nginx và PostgreSQL.')

# Page 13
page('3.3 Truy vấn LogQL')
p('Ba truy vấn dưới đây đã được chạy qua Loki API và giao diện Grafana Explore. Mỗi truy vấn giới hạn job/project, parse JSON rồi lọc trường cần quan sát. Khoảng thời gian Last 1 hour chứa cả thao tác demo và các kiểm thử có kiểm soát.')
h('Lỗi HTTP từ Nginx')
code('{job="nginx", project="de16hai"} | json | status >= 400')
p('Gửi POST không có CSRF tạo phản hồi 403 để kiểm chứng đường đi Nginx, app, Promtail và Loki.')
h('Vòng đời bảo trì')
code('{job="app", project="de16hai"} | json\n| event=~"maintenance_(scheduled|started|completed|cancelled)"')
figure('logql-maintenance.png','Hình 3.3 Explore hiển thị query và các sự kiện bảo trì đã lưu',6.1)
h('Cảnh báo và lỗi ứng dụng')
code('{job="app", project="de16hai"} | json | level=~"warn|error"')
p('Truy cập bộ lọc trạng thái không hợp lệ sau đăng nhập tạo request_rejected mức warn. check-logs.js xác nhận cả ba query có dòng phù hợp và stream PostgreSQL cũng đã được index. Promtail được giữ theo yêu cầu môn học dù đã EOL từ 02/03/2026; hướng nâng cấp là Grafana Alloy [4].')

# Page 14
page('CHƯƠNG 4 HARDENING VÀ ĐÁNH GIÁ')
h('4.1 Biện pháp gia cố')
p('Gia cố được thực hiện ở cả ứng dụng, container, mạng và database. Các biện pháp giảm quyền và diện tấn công cần kết hợp với kiểm thử, không bảo đảm an toàn tuyệt đối cho mọi tình huống.')
table(['Lớp bảo vệ','Cấu hình và bằng chứng'],[
('Quyền tiến trình','App UID 1000, Nginx UID 101; root filesystem chỉ đọc'),('Capability','App/Nginx bỏ ALL; bật no-new-privileges'),('Mạng và cổng','Database internal; app/DB/Loki/exporter không publish port; UI bind 127.0.0.1'),('Database','Tách bốn role; asset_app không superuser, không DELETE hồ sơ hoặc CREATE bảng'),('Bí mật','Secret ngoài Git, thư mục 0700; mật khẩu mạnh; scrypt lưu hash'),('HTTP và form','CSP, nosniff, chống iframe; CSRF; truy vấn SQL tham số hóa'),('Phiên và đăng nhập','Cookie HttpOnly/SameSite; giới hạn 15 lần/15 phút theo IP; auth_version'),('Tài nguyên','Giới hạn RAM/PID và thời gian statement/transaction DB')],[1.5,4.7])
p('check-hardening.py đọc Docker inspect để kiểm tra thực tế UID, filesystem, capabilities, network, port bindings và mount chỉ đọc của Promtail. Các kiểm thử HTTP/SQL xác nhận cookie, headers và quyền database, thay vì chỉ đọc cấu hình tĩnh.')
h('Ngoại lệ được giữ có chủ đích')
p('Storage init cần root cùng CHOWN/FOWNER/DAC_OVERRIDE để chuẩn bị volume rồi kết thúc. Logrotate chỉ có quyền cần thiết trên các volume log. PostgreSQL bootstrap bắt đầu root và hạ quyền sang postgres. cAdvisor chạy privileged, đọc host cgroup và có Docker socket; vì vậy đây là thành phần quản trị được tin cậy, không phải ranh giới cách ly an toàn cho mã không tin cậy.')
p('Demo chỉ mở localhost và dùng HTTP với security headers. Trước khi mở ra Internet cần TLS, cookie Secure, kiểm soát truy cập công cụ quản trị và rà soát lại các ngoại lệ quyền theo môi trường triển khai.')

# Page 15
page('4.2 Kết quả kiểm thử')
p('Bộ kiểm thử chạy qua Nginx và PostgreSQL thật trên Docker Engine trong Ubuntu. Tài khoản, thiết bị, danh mục và vị trí QA có dấu nhận diện riêng, được dọn sau mỗi lần chạy. Đổi mật khẩu chỉ thực hiện trên tài khoản QA, không thay tài khoản admin phục vụ demo.')
table(['Nhóm kiểm thử','Kết quả'],[
('Xác thực và CSRF','Khách bị chuyển đăng nhập; sai mật khẩu 401; POST rỗng/CSRF giả 403'),('Hồ sơ thiết bị','Tạo/sửa/tra cứu; mã/serial trùng bị từ chối; HTML được escape'),('Validation','Ngày, tiền, ID và khóa ngoại sai bị từ chối'),('Cập nhật đồng thời','Version cũ trả 409; hai start chỉ một lần thành công'),('Vòng đời bảo trì','Bắt đầu/hoàn tất/hủy đồng bộ trạng thái, giữ kết quả và chi phí'),('Phiên và quyền','Đổi mật khẩu vô hiệu phiên khác; app không có quyền xóa/tạo bảng'),('Hạ tầng','5 target UP; CPU/RAM có nhãn đúng; ba LogQL có kết quả'),('Lưu bền vững','Restart và seed lại giữ nguyên dữ liệu nghiệp vụ')],[1.7,4.5])
h('Kết quả Node Test Runner')
code('# tests 25\n# pass 25\n# fail 0\n# cancelled 0\n# skipped 0')
p('Kết quả được lưu ở .local/test-results/tests.tap. Tệp minh chứng công khai trong docs/evidence/verification.txt chỉ giữ thông tin kiểm tra đã lọc, không chứa secret hoặc dữ liệu phiên. Các script check-monitoring.js, check-logs.js và check-hardening.py bổ sung bằng chứng vận hành.')
p('Có thể chạy lại bằng sh scripts/test.sh. Khi vừa khởi động hoặc tạo lại app, cần đủ hai chu kỳ scrape để biểu đồ rate và kiểm tra CPU có dữ liệu. Counter reset khi process chạy lại là hành vi dự kiến; dữ liệu tài sản và lịch sử bảo trì nằm trong PostgreSQL vẫn được giữ.')

# Page 16
page('4.3 Sao lưu và phục hồi')
p('backup.sh dùng pg_dump dạng custom, tạo file và checksum SHA-256 trong .local/backups. Archive chứa schema, dữ liệu, owner và ACL, bao gồm tài khoản/phiên nên không được đưa lên GitHub. File backup có quyền 0600; thư mục riêng dùng 0700.')
code('sh scripts/backup.sh\nsh scripts/restore.sh --verify-current <duong-dan-backup>')
p('restore.sh kiểm tra checksum nếu có rồi tạo database tạm, phục hồi archive và tính fingerprint của bốn bảng nghiệp vụ. Chế độ --verify-current yêu cầu fingerprint và số dòng khớp database hiện tại; --verify cho phép bản cũ khác dữ liệu hiện tại nhưng vẫn kiểm tra phục hồi hợp lệ. Database tạm được xóa sau khi kiểm tra.')
table(['Nội dung kiểm chứng','Kết quả thực tế'],[
('Thiết bị phục hồi','12 bản ghi'),('Phiếu bảo trì phục hồi','3 bản ghi'),('Đối chiếu nội dung nghiệp vụ','Fingerprint khớp database hiện tại'),('Archive bị thay một byte','SHA-256 từ chối trước khi phục hồi'),('Sau restart và seed lại','Nội dung vẫn khớp bản đã sao lưu')],[2.4,3.8])
code('PASS: restored business rows match current database exactly.\nPASS: corrupt backup rejected by SHA-256 before restore.')
p('Bài thực hành kiểm chứng bằng database tạm, giữ database demo đang chạy. Khi chủ động cần thay dữ liệu hiện có, --apply kiểm tra archive trước, yêu cầu nhập PHUC HOI HAI và dừng app/Nginx trong lúc phục hồi. Restore áp dụng cả owner/ACL đã lưu; nếu thất bại, app được giữ dừng để xử lý.')
p('Bản sao lưu hiện nằm cùng máy nên chưa bảo vệ trước mất toàn bộ máy/ổ đĩa. Hướng vận hành tiếp theo là lưu thêm bản mã hóa ngoài máy và lập lịch thử phục hồi định kỳ. Mật khẩu người dùng sau restore theo bản sao lưu, không tự đồng bộ lại các file mật khẩu khởi tạo.')

# Page 17
page('KẾT LUẬN VÀ KỊCH BẢN DEMO')
p('Hải Asset đáp ứng phạm vi Đề 16: quản lý thiết bị và lịch sử bảo trì, lưu PostgreSQL, dùng pgAdmin, truy cập qua Nginx và triển khai bằng Docker Compose. Môi trường chạy là Docker Engine trong Ubuntu WSL2. Giám sát có số liệu container/web/DB; Loki truy vấn được log từ ba nguồn; các biện pháp hardening và phục hồi đã được kiểm chứng.')
p('Kết quả kiểm thử gồm 25 test đạt, các kiểm tra vận hành thành công và phục hồi 12 thiết bị/3 phiếu vào database tạm với dữ liệu khớp. Ảnh và sơ đồ được lưu cùng báo cáo, nguồn DOT/SVG cho phép chỉnh sửa khi cấu hình thay đổi. Mã nguồn và tài liệu được quản lý theo năm mốc commit có nội dung rõ ràng.')
h('Trình tự trình bày')
for item in [
 'Mở Ubuntu, kiểm tra Docker Engine và chạy start.sh. Mở website, pgAdmin, Grafana và Prometheus từ Chrome Windows.',
 'Đăng nhập website; tạo thiết bị có mã mới, kiểm tra thông tin và thử tìm kiếm/bộ lọc.',
 'Lập lịch, bắt đầu và hoàn tất bảo trì. Đối chiếu trạng thái trước/sau, kết quả và chi phí trong chi tiết thiết bị.',
 'Minh họa dữ liệu bị từ chối; chạy test.sh và check-hardening.py để chỉ ra kiểm soát đồng thời, phiên và quyền.',
 'Chỉ ra PostgreSQL trong pgAdmin; xem năm target UP và dashboard có CPU/RAM, HTTP, Nginx, PostgreSQL.',
 'Mở Explore với Loki, chạy ba LogQL và tìm request_id hoặc phiếu từ thao tác vừa làm.',
 'Tạo backup và chạy verify-current vào database tạm; giải thích cách dừng hệ thống mà giữ dữ liệu.'
]:
    q=doc.add_paragraph(item,'List Number');q.paragraph_format.space_after=Pt(6)
p('Hướng phát triển phù hợp gồm TLS khi triển khai ngoài localhost, phân quyền chi tiết theo vai trò, quản lý cấp phát/thu hồi và chuyển Promtail sang Grafana Alloy. Các tính năng này chưa thuộc phạm vi nghiệm thu hiện tại.')

# Page 18
page('TÀI LIỆU THAM KHẢO')
references=[
('[1]','Đề tài thực hành môn Triển khai và Quản trị Hệ thống Phần mềm K23','Tài liệu giảng viên cung cấp; Đề 16 và yêu cầu chung, trang 1, 2, 5, 6.'),
('[2]','Docker Engine trên Ubuntu','https://docs.docker.com/engine/install/ubuntu/'),
('[3]','Microsoft WSL và systemd','https://learn.microsoft.com/windows/wsl/systemd'),
('[4]','Grafana Loki Promtail và thông báo hết hỗ trợ','https://grafana.com/docs/loki/latest/send-data/promtail/'),
('[5]','Prometheus cấu hình scrape và cảnh báo','https://prometheus.io/docs/prometheus/latest/configuration/configuration/'),
('[6]','PostgreSQL pg_dump và pg_restore','https://www.postgresql.org/docs/17/app-pgdump.html\nhttps://www.postgresql.org/docs/17/app-pgrestore.html'),
('[7]','Grafana provisioning datasource và dashboard','https://grafana.com/docs/grafana/latest/administration/provisioning/'),
('[8]','Node.js LTS và Express','https://nodejs.org/en/about/previous-releases\nhttps://expressjs.com/'),
('[9]','Mã nguồn và cấu hình Hải Asset','https://github.com/dtc245200551-jpg/dtc245200551-de16-quan-ly-tai-san')]
for key,title,url in references:
    q=p(key+' '+title);q.runs[0].bold=True
    q=p(url);q.paragraph_format.space_after=Pt(13)
    for r in q.runs:r.font.size=Pt(11)
p('Các đường dẫn kỹ thuật là tài liệu chính thức của nhà cung cấp. Cấu hình và phiên bản thực tế được khóa trong repository; ảnh và kết quả kiểm tra phản ánh môi trường Ubuntu WSL2 của bài làm trong tháng 10 năm 2026.')

OUTPUT=ROOT/'docs/Bao_Cao_De16_Nguyen_Xuan_Hai.docx'
doc.save(OUTPUT)
print(OUTPUT)
print('Planned pages: 18; paragraphs:',len(doc.paragraphs),'tables:',len(doc.tables),'figures:',len(doc.inline_shapes))
